import io
import os
import re
import tempfile
import zipfile
import logging
from typing import List, Optional, Tuple

import pymupdf
from docx import Document
from fastapi import HTTPException, UploadFile, status

from app.core.config import settings
from app.models.parser import ParsedResumeResponse
from app.models.resume import ContactInfo

logger = logging.getLogger("resume_parser")

# Magic bytes
PDF_MAGIC_BYTES = b"%PDF-"
ZIP_MAGIC_BYTES = b"PK\x03\x04"

# Permitted MIME types
ALLOWED_MIME_TYPES = {
    ".pdf": [
        "application/pdf",
        "application/x-pdf",
        "application/acrobat",
        "applications/vnd.pdf",
        "text/pdf",
        "text/x-pdf",
        "application/octet-stream",  # Frequently dispatched by browsers
    ],
    ".docx": [
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "application/msword",
        "application/zip",
        "application/octet-stream",
    ],
}


class ResumeParsingError(HTTPException):
    """Custom HTTP Exception for resume parsing."""
    def __init__(self, status_code: int, detail: str):
        super().__init__(status_code=status_code, detail=detail)


class ResumeParserService:
    @staticmethod
    def normalize_text(text: str) -> str:
        """
        Normalize whitespace while preserving headings, paragraphs, and list bullets.
        Filters non-printable ASCII and control characters safely.
        """
        if not text:
            return ""

        # Remove null bytes and non-printable control characters (except newline, tab, carriage return)
        cleaned = re.sub(r"[\x00-\x08\x0B\x0C\x0E-\x1F\x7F]", "", text)
        # Normalize newlines
        cleaned = cleaned.replace("\r\n", "\n").replace("\r", "\n")

        # Process line by line to preserve bullets and headings
        lines = []
        for line in cleaned.split("\n"):
            # Collapse multiple inline spaces and tabs
            norm_line = re.sub(r"[ \t]+", " ", line).strip()
            lines.append(norm_line)

        reconstructed = "\n".join(lines)
        # Collapse 3 or more consecutive newlines into 2 (preserving section breaks)
        reconstructed = re.sub(r"\n{3,}", "\n\n", reconstructed)
        return reconstructed.strip()

    @classmethod
    def validate_file(cls, filename: str, content_type: Optional[str], file_bytes: bytes) -> str:
        """
        Validates extension, MIME type, file size, and magic bytes.
        Returns the normalized file extension.
        """
        # 1. Filename & Extension validation
        if not filename:
            raise ResumeParsingError(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Filename is missing or invalid.",
            )

        ext = os.path.splitext(filename.lower())[1]
        if ext not in settings.ALLOWED_EXTENSIONS:
            raise ResumeParsingError(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Unsupported file format '{ext}'. Only PDF (.pdf) and Word (.docx) files are supported.",
            )

        # 2. Empty file check
        file_size = len(file_bytes)
        if file_size == 0:
            raise ResumeParsingError(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Uploaded file is empty (0 bytes). Please upload a valid resume document.",
            )

        # 3. Size limit check
        if file_size > settings.max_file_size_bytes:
            raise ResumeParsingError(
                status_code=status.HTTP_413_CONTENT_TOO_LARGE,
                detail=f"File exceeds maximum allowed upload size of {settings.MAX_FILE_SIZE_MB}MB.",
            )

        # 4. MIME type validation if provided
        if content_type:
            allowed_mimes = ALLOWED_MIME_TYPES.get(ext, [])
            clean_mime = content_type.split(";")[0].strip().lower()
            if clean_mime not in allowed_mimes:
                logger.warning(f"Unusual MIME type '{clean_mime}' for extension '{ext}'")

        # 5. Magic Byte Inspection
        if ext == ".pdf":
            # PDF magic number %PDF- should appear in the first 1024 bytes
            if PDF_MAGIC_BYTES not in file_bytes[:1024]:
                raise ResumeParsingError(
                    status_code=status.HTTP_422_UNPROCESSABLE_CONTENT,
                    detail="Malformed document: File has a .pdf extension but lacks valid PDF signature bytes.",
                )
        elif ext == ".docx":
            # DOCX must be a valid zip archive
            if not file_bytes.startswith(ZIP_MAGIC_BYTES):
                raise ResumeParsingError(
                    status_code=status.HTTP_422_UNPROCESSABLE_CONTENT,
                    detail="Malformed document: File has a .docx extension but is not a valid OpenXML package.",
                )

        return ext

    @classmethod
    def extract_pdf_with_pymupdf(cls, file_bytes: bytes) -> Tuple[str, int, List[str]]:
        """
        Extract text and layout from PDF using PyMuPDF.
        Detects encryption, corrupt structures, and image-only (scanned) pages.
        """
        temp_path = None
        warnings = []
        try:
            with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as tmp:
                tmp.write(file_bytes)
                temp_path = tmp.name

            doc = pymupdf.open(temp_path)

            if doc.is_encrypted:
                raise ResumeParsingError(
                    status_code=status.HTTP_422_UNPROCESSABLE_CONTENT,
                    detail="The PDF document is password-protected or encrypted. Please upload an unencrypted file.",
                )

            page_count = len(doc)
            extracted_pages = []

            for page_num in range(page_count):
                page = doc[page_num]
                # Extract text preserving blocks and reading order
                text = page.get_text("text") or ""
                if text.strip():
                    extracted_pages.append(text.strip())

            doc.close()

            full_text = "\n\n".join(extracted_pages)
            normalized = cls.normalize_text(full_text)

            # Check for scanned or image-only PDF
            if not normalized or len(normalized.strip()) < 30:
                raise ResumeParsingError(
                    status_code=status.HTTP_422_UNPROCESSABLE_CONTENT,
                    detail="Document contains no extractable text. Scanned or image-only documents require OCR (optical character recognition) which is currently not supported. Please upload a text-based PDF or DOCX file.",
                )

            return normalized, page_count, warnings

        except ResumeParsingError:
            raise
        except Exception as e:
            logger.error(f"Error parsing PDF: {type(e).__name__}")
            raise ResumeParsingError(
                status_code=status.HTTP_422_UNPROCESSABLE_CONTENT,
                detail=f"Failed to parse PDF document: The file appears to be corrupted or invalid.",
            )
        finally:
            if temp_path and os.path.exists(temp_path):
                try:
                    os.unlink(temp_path)
                except Exception:
                    pass

    @classmethod
    def extract_docx(cls, file_bytes: bytes) -> Tuple[str, Optional[int], List[str]]:
        """
        Safely extract text, paragraphs, tables, and bullets from DOCX bytes.
        Verifies internal XML structure.
        """
        temp_path = None
        warnings = []
        try:
            # Check internal zip contents
            try:
                with zipfile.ZipFile(io.BytesIO(file_bytes)) as z:
                    namelist = z.namelist()
                    if "word/document.xml" not in namelist:
                        raise ResumeParsingError(
                            status_code=status.HTTP_422_UNPROCESSABLE_CONTENT,
                            detail="Malformed document: DOCX package missing word/document.xml entry.",
                        )
            except zipfile.BadZipFile:
                raise ResumeParsingError(
                    status_code=status.HTTP_422_UNPROCESSABLE_CONTENT,
                    detail="Malformed document: Corrupt or damaged DOCX zip archive.",
                )

            with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tmp:
                tmp.write(file_bytes)
                temp_path = tmp.name

            doc = Document(temp_path)
            lines = []

            # Extract paragraphs
            for p in doc.paragraphs:
                p_text = p.text.strip()
                if p_text:
                    # Check if style is bullet
                    if "bullet" in p.style.name.lower():
                        lines.append(f"• {p_text}")
                    else:
                        lines.append(p_text)

            # Extract tables
            for table in doc.tables:
                for row in table.rows:
                    row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_cells:
                        lines.append(" | ".join(row_cells))

            raw_text = "\n".join(lines)
            normalized = cls.normalize_text(raw_text)

            if not normalized or len(normalized.strip()) < 30:
                raise ResumeParsingError(
                    status_code=status.HTTP_422_UNPROCESSABLE_CONTENT,
                    detail="Document contains no extractable text. Please ensure your DOCX file contains readable text.",
                )

            return normalized, None, warnings

        except ResumeParsingError:
            raise
        except Exception as e:
            logger.error(f"Error parsing DOCX: {type(e).__name__}")
            raise ResumeParsingError(
                status_code=status.HTTP_422_UNPROCESSABLE_CONTENT,
                detail="Failed to parse DOCX document: The file appears to be corrupted or invalid.",
            )
        finally:
            if temp_path and os.path.exists(temp_path):
                try:
                    os.unlink(temp_path)
                except Exception:
                    pass

    @classmethod
    def extract_contact_info(cls, text: str) -> ContactInfo:
        """Extract contact heuristics safely."""
        email_match = re.search(r"[\w\.-]+@[\w\.-]+\.\w+", text)
        phone_match = re.search(r"(?:(?:\+?\d{1,3}[-.\s]?)?\(?\d{3}\)?[-.\s]?\d{3}[-.\s]?\d{4})", text)
        linkedin_match = re.search(r"(?:https?:\/\/)?(?:www\.)?linkedin\.com\/in\/[\w\-]+", text, re.IGNORECASE)
        github_match = re.search(r"(?:https?:\/\/)?(?:www\.)?github\.com\/[\w\-]+", text, re.IGNORECASE)

        first_line = ""
        for line in text.split("\n"):
            line = line.strip()
            if line and len(line) < 50 and not any(k in line.lower() for k in ["@", "http", "phone", "resume", "curriculum", "page"]):
                first_line = line
                break

        return ContactInfo(
            name=first_line,
            email=email_match.group(0) if email_match else "",
            phone=phone_match.group(0) if phone_match else "",
            linkedin=linkedin_match.group(0) if linkedin_match else None,
            github=github_match.group(0) if github_match else None,
        )

    @classmethod
    def parse_upload(cls, file: UploadFile) -> ParsedResumeResponse:
        """
        Orchestrates full file parsing: validation, extraction, normalization, and metadata generation.
        """
        # Read content safely
        try:
            content = file.file.read()
        except Exception as e:
            raise ResumeParsingError(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Could not read uploaded stream: {str(e)}",
            )

        filename = file.filename or "uploaded_file"
        ext = cls.validate_file(filename, file.content_type, content)

        if ext == ".pdf":
            text, page_count, warnings = cls.extract_pdf_with_pymupdf(content)
            file_type = "pdf"
        elif ext == ".docx":
            text, page_count, warnings = cls.extract_docx(content)
            file_type = "docx"
        else:
            raise ResumeParsingError(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Unsupported format.",
            )

        contact_info = cls.extract_contact_info(text)
        word_count = len(text.split())
        char_count = len(text)
        is_scanned = char_count < 100

        if is_scanned:
            warnings.append("Document has unusually low text density; verify that content was not scanned as an image.")

        return ParsedResumeResponse(
            filename=filename,
            file_type=file_type,
            file_size_bytes=len(content),
            page_count=page_count,
            word_count=word_count,
            character_count=char_count,
            text=text,
            contact_info=contact_info,
            is_scanned_or_low_text=is_scanned,
            warnings=warnings,
        )


resume_parser_service = ResumeParserService()
