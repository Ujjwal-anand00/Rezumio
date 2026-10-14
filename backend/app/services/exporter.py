import io
import re
from typing import Any, Dict, List, Union
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, HRFlowable
from reportlab.lib import colors
from reportlab.lib.units import inch


from app.models.resume import ContactInfo


class DocumentExporterService:
    @staticmethod
    def _extract_fields(resume: Any) -> Dict[str, Any]:
        """Normalize attributes from either TailoredResumeResponse or CustomizeResumeResponse or dict."""
        if isinstance(resume, dict):
            contact = resume.get("contact_info", {})
            if isinstance(contact, dict):
                contact_info = ContactInfo(**contact)
            else:
                contact_info = contact

            summary = resume.get("professional_summary") or resume.get("tailored_summary") or ""
            exp = resume.get("experience") or resume.get("work_experience") or []
            skills = resume.get("relevant_existing_skills") or resume.get("skills") or {}
            projects = resume.get("projects") or []
            education = resume.get("education") or []
            certs = resume.get("certifications") or []
        else:
            contact_info = getattr(resume, "contact_info", ContactInfo())
            summary = getattr(resume, "professional_summary", None) or getattr(resume, "tailored_summary", "")
            exp = getattr(resume, "experience", None) or getattr(resume, "work_experience", [])
            skills = getattr(resume, "relevant_existing_skills", None) or getattr(resume, "skills", {})
            projects = getattr(resume, "projects", [])
            education = getattr(resume, "education", [])
            certs = getattr(resume, "certifications", [])

        return {
            "contact_info": contact_info,
            "summary": summary,
            "experience": exp,
            "skills": skills,
            "projects": projects,
            "education": education,
            "certifications": certs,
        }

    @classmethod
    def generate_docx(cls, resume: Any) -> io.BytesIO:
        data = cls._extract_fields(resume)
        contact: ContactInfo = data["contact_info"]

        doc = Document()

        # Set standard 0.75 inch ATS margins
        for section in doc.sections:
            section.top_margin = Inches(0.75)
            section.bottom_margin = Inches(0.75)
            section.left_margin = Inches(0.75)
            section.right_margin = Inches(0.75)

        # Name / Header
        name_p = doc.add_paragraph()
        name_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        name_run = name_p.add_run(getattr(contact, "name", "Candidate Name") or "Candidate Name")
        name_run.font.size = Pt(20)
        name_run.font.bold = True
        name_run.font.name = "Arial"
        name_run.font.color.rgb = RGBColor(15, 23, 42)

        # Contact Info Line
        contact_parts = []
        if getattr(contact, "email", None):
            contact_parts.append(contact.email)
        if getattr(contact, "phone", None):
            contact_parts.append(contact.phone)
        if getattr(contact, "location", None):
            contact_parts.append(contact.location)
        if getattr(contact, "linkedin", None):
            contact_parts.append(contact.linkedin)
        if getattr(contact, "github", None):
            contact_parts.append(contact.github)

        if contact_parts:
            contact_p = doc.add_paragraph()
            contact_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            contact_run = contact_p.add_run(" | ".join(contact_parts))
            contact_run.font.size = Pt(9.5)
            contact_run.font.name = "Arial"
            contact_run.font.color.rgb = RGBColor(100, 116, 139)

        doc.add_paragraph()

        # Section Header Helper
        def add_section_header(title: str):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(3)
            run = p.add_run(title.upper())
            run.font.size = Pt(11)
            run.font.bold = True
            run.font.name = "Arial"
            run.font.color.rgb = RGBColor(15, 23, 42)

        # Professional Summary
        if data["summary"]:
            add_section_header("Professional Summary")
            sp = doc.add_paragraph()
            sp.paragraph_format.space_after = Pt(6)
            srun = sp.add_run(data["summary"])
            srun.font.size = Pt(10)
            srun.font.name = "Arial"

        # Work Experience
        if data["experience"]:
            add_section_header("Work Experience")
            for item in data["experience"]:
                head_p = doc.add_paragraph()
                head_p.paragraph_format.space_before = Pt(5)
                head_p.paragraph_format.space_after = Pt(1)

                position = getattr(item, "position", "") if not isinstance(item, dict) else item.get("position", "")
                company = getattr(item, "company", "") if not isinstance(item, dict) else item.get("company", "")
                location = getattr(item, "location", "") if not isinstance(item, dict) else item.get("location", "")
                start_date = getattr(item, "start_date", "") if not isinstance(item, dict) else item.get("start_date", "")
                end_date = getattr(item, "end_date", "") if not isinstance(item, dict) else item.get("end_date", "")
                is_current = getattr(item, "is_current", False) if not isinstance(item, dict) else item.get("is_current", False)

                role_run = head_p.add_run(position)
                role_run.bold = True
                role_run.font.size = Pt(10.5)
                role_run.font.name = "Arial"

                comp_text = f" | {company}"
                if location:
                    comp_text += f", {location}"
                dates_text = f" ({start_date or ''} – {'Present' if is_current else end_date or 'Present'})"
                comp_text += dates_text

                comp_run = head_p.add_run(comp_text)
                comp_run.font.size = Pt(10)
                comp_run.font.name = "Arial"
                comp_run.font.color.rgb = RGBColor(71, 85, 105)

                bullets = getattr(item, "bullets", []) if not isinstance(item, dict) else item.get("bullets", [])
                for b in bullets:
                    b_text = ""
                    if isinstance(b, str):
                        b_text = b
                    elif isinstance(b, dict):
                        b_text = b.get("tailored") or b.get("original") or ""
                    else:
                        b_text = getattr(b, "tailored", None) or getattr(b, "original", "")

                    if b_text:
                        bp = doc.add_paragraph(style="List Bullet")
                        bp.paragraph_format.space_before = Pt(1)
                        bp.paragraph_format.space_after = Pt(2)
                        brun = bp.add_run(b_text)
                        brun.font.size = Pt(9.5)
                        brun.font.name = "Arial"

        # Technical Skills
        if data["skills"]:
            add_section_header("Technical Skills")
            for cat, items in data["skills"].items():
                if items:
                    kp = doc.add_paragraph()
                    kp.paragraph_format.space_after = Pt(2)
                    c_run = kp.add_run(f"{cat}: ")
                    c_run.bold = True
                    c_run.font.size = Pt(9.5)
                    c_run.font.name = "Arial"
                    i_run = kp.add_run(", ".join(items))
                    i_run.font.size = Pt(9.5)
                    i_run.font.name = "Arial"

        # Projects
        if data["projects"]:
            add_section_header("Key Projects")
            for proj in data["projects"]:
                pp = doc.add_paragraph()
                pp.paragraph_format.space_before = Pt(3)
                pp.paragraph_format.space_after = Pt(1)

                title = getattr(proj, "title", "") if not isinstance(proj, dict) else proj.get("title", "")
                techs = getattr(proj, "technologies", []) if not isinstance(proj, dict) else proj.get("technologies", [])
                p_run = pp.add_run(title)
                p_run.bold = True
                p_run.font.size = Pt(10)
                p_run.font.name = "Arial"

                if techs:
                    t_run = pp.add_run(f" ({', '.join(techs)})")
                    t_run.font.size = Pt(9.5)
                    t_run.font.italic = True
                    t_run.font.name = "Arial"

                p_bullets = getattr(proj, "bullet_points", []) if not isinstance(proj, dict) else proj.get("bullet_points", [])
                for pb in p_bullets:
                    pbp = doc.add_paragraph(style="List Bullet")
                    pbp.paragraph_format.space_before = Pt(1)
                    pbp.paragraph_format.space_after = Pt(1)
                    prun = pbp.add_run(pb)
                    prun.font.size = Pt(9.5)
                    prun.font.name = "Arial"

        # Education
        if data["education"]:
            add_section_header("Education")
            for edu in data["education"]:
                ep = doc.add_paragraph()
                ep.paragraph_format.space_after = Pt(2)
                degree = getattr(edu, "degree", "") if not isinstance(edu, dict) else edu.get("degree", "")
                institution = getattr(edu, "institution", "") if not isinstance(edu, dict) else edu.get("institution", "")
                grad = getattr(edu, "graduation_date", "") if not isinstance(edu, dict) else edu.get("graduation_date", "")
                gpa = getattr(edu, "gpa", "") if not isinstance(edu, dict) else edu.get("gpa", "")

                deg_run = ep.add_run(degree)
                deg_run.bold = True
                deg_run.font.size = Pt(10)
                deg_run.font.name = "Arial"

                detail_text = f" – {institution}"
                if grad:
                    detail_text += f" ({grad})"
                if gpa:
                    detail_text += f" | GPA: {gpa}"
                d_run = ep.add_run(detail_text)
                d_run.font.size = Pt(9.5)
                d_run.font.name = "Arial"

        # Certifications
        if data["certifications"]:
            add_section_header("Certifications")
            for cert in data["certifications"]:
                cp = doc.add_paragraph(style="List Bullet")
                cp.paragraph_format.space_after = Pt(1)
                c_run = cp.add_run(cert)
                c_run.font.size = Pt(9.5)
                c_run.font.name = "Arial"

        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        return buffer

    @classmethod
    def generate_pdf(cls, resume: Any) -> io.BytesIO:
        data = cls._extract_fields(resume)
        contact: ContactInfo = data["contact_info"]

        buffer = io.BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=letter,
            rightMargin=54,  # 0.75 in
            leftMargin=54,
            topMargin=54,
            bottomMargin=54,
        )

        styles = getSampleStyleSheet()

        title_style = ParagraphStyle(
            "ResumeTitle",
            parent=styles["Normal"],
            fontName="Helvetica-Bold",
            fontSize=18,
            leading=22,
            textColor=colors.HexColor("#0f172a"),
            alignment=1,  # Center
        )

        contact_style = ParagraphStyle(
            "ResumeContact",
            parent=styles["Normal"],
            fontName="Helvetica",
            fontSize=9,
            leading=12,
            textColor=colors.HexColor("#475569"),
            alignment=1,
        )

        heading_style = ParagraphStyle(
            "ResumeHeading",
            parent=styles["Normal"],
            fontName="Helvetica-Bold",
            fontSize=11,
            leading=14,
            textColor=colors.HexColor("#1e293b"),
            spaceBefore=8,
            spaceAfter=2,
        )

        role_style = ParagraphStyle(
            "ResumeRole",
            parent=styles["Normal"],
            fontName="Helvetica-Bold",
            fontSize=10,
            leading=13,
            textColor=colors.HexColor("#0f172a"),
        )

        bullet_style = ParagraphStyle(
            "ResumeBullet",
            parent=styles["Normal"],
            fontName="Helvetica",
            fontSize=9.5,
            leading=13,
            leftIndent=14,
            firstLineIndent=-10,
            textColor=colors.HexColor("#334155"),
            spaceAfter=2,
        )

        body_style = ParagraphStyle(
            "ResumeBody",
            parent=styles["Normal"],
            fontName="Helvetica",
            fontSize=9.5,
            leading=13,
            textColor=colors.HexColor("#334155"),
            spaceAfter=4,
        )

        elements = []

        # Name
        elements.append(Paragraph(getattr(contact, "name", "Candidate Name") or "Candidate Name", title_style))
        elements.append(Spacer(1, 4))

        # Contact line
        contact_parts = []
        if getattr(contact, "email", None):
            contact_parts.append(contact.email)
        if getattr(contact, "phone", None):
            contact_parts.append(contact.phone)
        if getattr(contact, "location", None):
            contact_parts.append(contact.location)
        if getattr(contact, "linkedin", None):
            contact_parts.append(contact.linkedin)
        if getattr(contact, "github", None):
            contact_parts.append(contact.github)

        if contact_parts:
            elements.append(Paragraph(" &bull; ".join(contact_parts), contact_style))
            elements.append(Spacer(1, 8))

        def add_pdf_divider():
            elements.append(HRFlowable(width="100%", thickness=0.75, color=colors.HexColor("#cbd5e1"), spaceBefore=2, spaceAfter=4))

        # Summary
        if data["summary"]:
            elements.append(Paragraph("PROFESSIONAL SUMMARY", heading_style))
            add_pdf_divider()
            elements.append(Paragraph(data["summary"], body_style))
            elements.append(Spacer(1, 4))

        # Experience
        if data["experience"]:
            elements.append(Paragraph("WORK EXPERIENCE", heading_style))
            add_pdf_divider()
            for exp in data["experience"]:
                pos = getattr(exp, "position", "") if not isinstance(exp, dict) else exp.get("position", "")
                comp = getattr(exp, "company", "") if not isinstance(exp, dict) else exp.get("company", "")
                loc = getattr(exp, "location", "") if not isinstance(exp, dict) else exp.get("location", "")
                s_date = getattr(exp, "start_date", "") if not isinstance(exp, dict) else exp.get("start_date", "")
                e_date = getattr(exp, "end_date", "") if not isinstance(exp, dict) else exp.get("end_date", "")
                cur = getattr(exp, "is_current", False) if not isinstance(exp, dict) else exp.get("is_current", False)

                date_str = f"{s_date or ''} - {'Present' if cur else e_date or 'Present'}"
                company_str = f"<b>{pos}</b> | {comp}"
                if loc:
                    company_str += f", {loc}"
                company_str += f" <font color='#64748b'>({date_str})</font>"

                elements.append(Paragraph(company_str, role_style))
                elements.append(Spacer(1, 2))

                bullets = getattr(exp, "bullets", []) if not isinstance(exp, dict) else exp.get("bullets", [])
                for b in bullets:
                    b_text = ""
                    if isinstance(b, str):
                        b_text = b
                    elif isinstance(b, dict):
                        b_text = b.get("tailored") or b.get("original") or ""
                    else:
                        b_text = getattr(b, "tailored", None) or getattr(b, "original", "")

                    if b_text:
                        elements.append(Paragraph(f"&bull; {b_text}", bullet_style))

                elements.append(Spacer(1, 4))

        # Skills
        if data["skills"]:
            elements.append(Paragraph("TECHNICAL SKILLS", heading_style))
            add_pdf_divider()
            for cat, items in data["skills"].items():
                if items:
                    line = f"<b>{cat}:</b> {', '.join(items)}"
                    elements.append(Paragraph(line, body_style))
            elements.append(Spacer(1, 4))

        # Projects
        if data["projects"]:
            elements.append(Paragraph("KEY PROJECTS", heading_style))
            add_pdf_divider()
            for proj in data["projects"]:
                title = getattr(proj, "title", "") if not isinstance(proj, dict) else proj.get("title", "")
                techs = getattr(proj, "technologies", []) if not isinstance(proj, dict) else proj.get("technologies", [])
                p_line = f"<b>{title}</b>"
                if techs:
                    p_line += f" <font color='#64748b'>({', '.join(techs)})</font>"
                elements.append(Paragraph(p_line, role_style))
                elements.append(Spacer(1, 1))

                p_bullets = getattr(proj, "bullet_points", []) if not isinstance(proj, dict) else proj.get("bullet_points", [])
                for pb in p_bullets:
                    elements.append(Paragraph(f"&bull; {pb}", bullet_style))
                elements.append(Spacer(1, 3))

        # Education
        if data["education"]:
            elements.append(Paragraph("EDUCATION", heading_style))
            add_pdf_divider()
            for edu in data["education"]:
                degree = getattr(edu, "degree", "") if not isinstance(edu, dict) else edu.get("degree", "")
                inst = getattr(edu, "institution", "") if not isinstance(edu, dict) else edu.get("institution", "")
                grad = getattr(edu, "graduation_date", "") if not isinstance(edu, dict) else edu.get("graduation_date", "")
                gpa = getattr(edu, "gpa", "") if not isinstance(edu, dict) else edu.get("gpa", "")

                line = f"<b>{degree}</b> - {inst}"
                if grad:
                    line += f" <font color='#64748b'>({grad})</font>"
                if gpa:
                    line += f" | GPA: {gpa}"
                elements.append(Paragraph(line, body_style))
            elements.append(Spacer(1, 4))

        # Certifications
        if data["certifications"]:
            elements.append(Paragraph("CERTIFICATIONS", heading_style))
            add_pdf_divider()
            for cert in data["certifications"]:
                elements.append(Paragraph(f"&bull; {cert}", bullet_style))

        doc.build(elements)
        buffer.seek(0)
        return buffer

    @classmethod
    def generate_txt(cls, resume: Any) -> str:
        data = cls._extract_fields(resume)
        contact: ContactInfo = data["contact_info"]

        lines = []
        name = getattr(contact, "name", "Candidate Name") or "Candidate Name"
        lines.append(name.upper())
        lines.append("=" * len(name))

        contact_parts = []
        if getattr(contact, "email", None):
            contact_parts.append(contact.email)
        if getattr(contact, "phone", None):
            contact_parts.append(contact.phone)
        if getattr(contact, "location", None):
            contact_parts.append(contact.location)
        if getattr(contact, "linkedin", None):
            contact_parts.append(contact.linkedin)
        if getattr(contact, "github", None):
            contact_parts.append(contact.github)

        if contact_parts:
            lines.append(" | ".join(contact_parts))
        lines.append("")

        if data["summary"]:
            lines.append("PROFESSIONAL SUMMARY")
            lines.append("-" * 20)
            lines.append(data["summary"])
            lines.append("")

        if data["experience"]:
            lines.append("WORK EXPERIENCE")
            lines.append("-" * 15)
            for exp in data["experience"]:
                pos = getattr(exp, "position", "") if not isinstance(exp, dict) else exp.get("position", "")
                comp = getattr(exp, "company", "") if not isinstance(exp, dict) else exp.get("company", "")
                s_date = getattr(exp, "start_date", "") if not isinstance(exp, dict) else exp.get("start_date", "")
                e_date = getattr(exp, "end_date", "") if not isinstance(exp, dict) else exp.get("end_date", "")
                cur = getattr(exp, "is_current", False) if not isinstance(exp, dict) else exp.get("is_current", False)

                date_str = f"({s_date or ''} - {'Present' if cur else e_date or 'Present'})"
                lines.append(f"{pos} | {comp} {date_str}")

                bullets = getattr(exp, "bullets", []) if not isinstance(exp, dict) else exp.get("bullets", [])
                for b in bullets:
                    b_text = ""
                    if isinstance(b, str):
                        b_text = b
                    elif isinstance(b, dict):
                        b_text = b.get("tailored") or b.get("original") or ""
                    else:
                        b_text = getattr(b, "tailored", None) or getattr(b, "original", "")

                    if b_text:
                        lines.append(f"  * {b_text}")
                lines.append("")

        if data["skills"]:
            lines.append("TECHNICAL SKILLS")
            lines.append("-" * 16)
            for cat, items in data["skills"].items():
                if items:
                    lines.append(f"  {cat}: {', '.join(items)}")
            lines.append("")

        if data["projects"]:
            lines.append("KEY PROJECTS")
            lines.append("-" * 12)
            for proj in data["projects"]:
                title = getattr(proj, "title", "") if not isinstance(proj, dict) else proj.get("title", "")
                lines.append(f"  {title}")
                p_bullets = getattr(proj, "bullet_points", []) if not isinstance(proj, dict) else proj.get("bullet_points", [])
                for pb in p_bullets:
                    lines.append(f"    * {pb}")
            lines.append("")

        if data["education"]:
            lines.append("EDUCATION")
            lines.append("-" * 9)
            for edu in data["education"]:
                degree = getattr(edu, "degree", "") if not isinstance(edu, dict) else edu.get("degree", "")
                inst = getattr(edu, "institution", "") if not isinstance(edu, dict) else edu.get("institution", "")
                grad = getattr(edu, "graduation_date", "") if not isinstance(edu, dict) else edu.get("graduation_date", "")
                lines.append(f"  {degree} - {inst} ({grad})")
            lines.append("")

        if data["certifications"]:
            lines.append("CERTIFICATIONS")
            lines.append("-" * 14)
            for cert in data["certifications"]:
                lines.append(f"  * {cert}")
            lines.append("")

        return "\n".join(lines)

    @classmethod
    def generate_cover_letter_txt(cls, cl: Any) -> str:
        name = getattr(cl, "candidate_name", "Candidate") if not isinstance(cl, dict) else cl.get("candidate_name", "Candidate")
        email = getattr(cl, "candidate_email", "") if not isinstance(cl, dict) else cl.get("candidate_email", "")
        phone = getattr(cl, "candidate_phone", "") if not isinstance(cl, dict) else cl.get("candidate_phone", "")
        company = getattr(cl, "company_name", "") if not isinstance(cl, dict) else cl.get("company_name", "")
        role = getattr(cl, "job_title", "") if not isinstance(cl, dict) else cl.get("job_title", "")
        date_str = getattr(cl, "date_formatted", "") if not isinstance(cl, dict) else cl.get("date_formatted", "")
        salutation = getattr(cl, "salutation", "Dear Hiring Team,") if not isinstance(cl, dict) else cl.get("salutation", "Dear Hiring Team,")
        paragraphs = getattr(cl, "paragraphs", []) if not isinstance(cl, dict) else cl.get("paragraphs", [])
        closing = getattr(cl, "closing_salutation", "Sincerely,") if not isinstance(cl, dict) else cl.get("closing_salutation", "Sincerely,")

        lines = [
            name.upper(),
            f"{email} | {phone}",
            "",
            date_str,
            "",
            f"RE: {role} Application - {company}",
            "",
            salutation,
            "",
        ]
        for p in paragraphs:
            content = getattr(p, "content", "") if not isinstance(p, dict) else p.get("content", "")
            lines.append(content)
            lines.append("")

        lines.append(closing)
        lines.append(name)
        return "\n".join(lines)

    @classmethod
    def generate_cover_letter_docx(cls, cl: Any) -> io.BytesIO:
        doc = Document()
        for s in doc.sections:
            s.top_margin = Inches(1.0)
            s.bottom_margin = Inches(1.0)
            s.left_margin = Inches(1.0)
            s.right_margin = Inches(1.0)

        name = getattr(cl, "candidate_name", "Candidate") if not isinstance(cl, dict) else cl.get("candidate_name", "Candidate")
        email = getattr(cl, "candidate_email", "") if not isinstance(cl, dict) else cl.get("candidate_email", "")
        phone = getattr(cl, "candidate_phone", "") if not isinstance(cl, dict) else cl.get("candidate_phone", "")
        company = getattr(cl, "company_name", "") if not isinstance(cl, dict) else cl.get("company_name", "")
        role = getattr(cl, "job_title", "") if not isinstance(cl, dict) else cl.get("job_title", "")
        date_str = getattr(cl, "date_formatted", "") if not isinstance(cl, dict) else cl.get("date_formatted", "")
        salutation = getattr(cl, "salutation", "Dear Hiring Team,") if not isinstance(cl, dict) else cl.get("salutation", "Dear Hiring Team,")
        paragraphs = getattr(cl, "paragraphs", []) if not isinstance(cl, dict) else cl.get("paragraphs", [])
        closing = getattr(cl, "closing_salutation", "Sincerely,") if not isinstance(cl, dict) else cl.get("closing_salutation", "Sincerely,")

        # Header
        h_p = doc.add_paragraph()
        run_name = h_p.add_run(name)
        run_name.font.name = "Calibri"
        run_name.font.size = Pt(16)
        run_name.font.bold = True
        run_name.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

        sub_p = doc.add_paragraph(f"{email} | {phone}")
        sub_p.paragraph_format.space_after = Pt(20)

        doc.add_paragraph(date_str)
        re_p = doc.add_paragraph(f"RE: {role} Application - {company}")
        re_p.runs[0].font.bold = True
        re_p.paragraph_format.space_after = Pt(12)

        doc.add_paragraph(salutation)

        for p in paragraphs:
            content = getattr(p, "content", "") if not isinstance(p, dict) else p.get("content", "")
            para = doc.add_paragraph(content)
            para.paragraph_format.space_after = Pt(10)
            para.paragraph_format.line_spacing = 1.15

        doc.add_paragraph(closing)
        sign_p = doc.add_paragraph(name)
        sign_p.runs[0].font.bold = True

        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        return buffer

    @classmethod
    def generate_cover_letter_pdf(cls, cl: Any) -> io.BytesIO:
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=letter,
            rightMargin=0.75 * inch,
            leftMargin=0.75 * inch,
            topMargin=0.75 * inch,
            bottomMargin=0.75 * inch,
        )

        name = getattr(cl, "candidate_name", "Candidate") if not isinstance(cl, dict) else cl.get("candidate_name", "Candidate")
        email = getattr(cl, "candidate_email", "") if not isinstance(cl, dict) else cl.get("candidate_email", "")
        phone = getattr(cl, "candidate_phone", "") if not isinstance(cl, dict) else cl.get("candidate_phone", "")
        company = getattr(cl, "company_name", "") if not isinstance(cl, dict) else cl.get("company_name", "")
        role = getattr(cl, "job_title", "") if not isinstance(cl, dict) else cl.get("job_title", "")
        date_str = getattr(cl, "date_formatted", "") if not isinstance(cl, dict) else cl.get("date_formatted", "")
        salutation = getattr(cl, "salutation", "Dear Hiring Team,") if not isinstance(cl, dict) else cl.get("salutation", "Dear Hiring Team,")
        paragraphs = getattr(cl, "paragraphs", []) if not isinstance(cl, dict) else cl.get("paragraphs", [])
        closing = getattr(cl, "closing_salutation", "Sincerely,") if not isinstance(cl, dict) else cl.get("closing_salutation", "Sincerely,")

        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            "CLTitle",
            parent=styles["Heading1"],
            fontSize=16,
            leading=20,
            textColor=colors.HexColor("#0f172a"),
            fontName="Helvetica-Bold",
            spaceAfter=4,
        )
        contact_style = ParagraphStyle(
            "CLContact",
            parent=styles["Normal"],
            fontSize=9.5,
            leading=12,
            textColor=colors.HexColor("#475569"),
            spaceAfter=16,
        )
        body_style = ParagraphStyle(
            "CLBody",
            parent=styles["Normal"],
            fontSize=10,
            leading=14.5,
            textColor=colors.HexColor("#1e293b"),
            spaceAfter=12,
        )
        bold_style = ParagraphStyle(
            "CLBold",
            parent=styles["Normal"],
            fontSize=10,
            leading=14,
            textColor=colors.HexColor("#0f172a"),
            fontName="Helvetica-Bold",
            spaceAfter=10,
        )

        story = [
            Paragraph(name, title_style),
            Paragraph(f"{email} &nbsp;|&nbsp; {phone}", contact_style),
            Paragraph(date_str, body_style),
            Paragraph(f"<b>RE: {role} Application - {company}</b>", bold_style),
            Paragraph(salutation, body_style),
        ]

        for p in paragraphs:
            content = getattr(p, "content", "") if not isinstance(p, dict) else p.get("content", "")
            story.append(Paragraph(content, body_style))

        story.append(Spacer(1, 10))
        story.append(Paragraph(closing, body_style))
        story.append(Spacer(1, 5))
        story.append(Paragraph(f"<b>{name}</b>", bold_style))

        doc.build(story)
        buffer.seek(0)
        return buffer


exporter_service = DocumentExporterService()

